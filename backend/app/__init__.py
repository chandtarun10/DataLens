import io
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PURPLE = '#6744ff'
NAVY = '#081b3d'
TEXT = '#1b2b4a'


def _pdf_table(data, widths=None):
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(NAVY)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#b9c6dc')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor('#f4f7fc')
        ]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5)
    ]))
    return t


def _pdf_value(value):
    if value is None:
        return ''
    if isinstance(value, float):
        return f'{value:.4g}'
    return str(value)


def _chart_png(title, kind, labels, values):
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    if kind == 'bar':
        labels = [str(label) for label in labels]
        positions = list(range(len(labels)))

        ax.bar(positions, values, color='#6744ff')

        # Keep category / date labels readable: rotate long or many
        # labels and show only every n-th one when there are many bars.
        step = max(1, int(np.ceil(len(labels) / 12)))
        longest = max((len(label) for label in labels), default=0)
        rotate = len(labels) > 6 or longest > 10

        ax.set_xticks(positions)
        ax.set_xticklabels(
            [
                label if i % step == 0 else ''
                for i, label in enumerate(labels)
            ],
            rotation=45 if rotate else 0,
            ha='right' if rotate else 'center',
            fontsize=8
        )

    elif kind == 'barh':
        labels = [str(label) for label in labels]

        ax.barh(labels, values, color='#6744ff')
        ax.invert_yaxis()
        ax.tick_params(axis='y', labelsize=8)

    elif kind == 'hist':
        values = np.asarray(values, dtype=float)

        if len(values):
            bins = min(
                15,
                max(5, int(np.sqrt(max(1, len(values)))))
            )
            ax.hist(values, bins=bins, color='#6744ff')

    elif kind == 'box':
        ax.boxplot(
            values,
            vert=True,
            patch_artist=True,
            boxprops={'facecolor': '#6744ff'}
        )

    ax.set_title(title)
    ax.grid(axis='x' if kind == 'barh' else 'y', alpha=.18)

    fig.tight_layout()

    bio = io.BytesIO()
    fig.savefig(
        bio,
        format='png',
        dpi=150,
        bbox_inches='tight'
    )

    plt.close(fig)
    bio.seek(0)

    return bio


def build_pdf(r, name):
    bio = io.BytesIO()

    doc = SimpleDocTemplate(
        bio,
        pagesize=A4,
        rightMargin=13 * mm,
        leftMargin=13 * mm,
        topMargin=13 * mm,
        bottomMargin=13 * mm
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name='Small',
            parent=styles['BodyText'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#40516d')
        )
    )

    # ---------------------------------------------------------
    # REPORT HEADER
    # ---------------------------------------------------------
    story = [
        Paragraph(
            'DataLens — Data Quality Analysis Report',
            styles['Title']
        ),
        Paragraph(
            f'File: {name}',
            styles['BodyText']
        ),
        Spacer(1, 8)
    ]

    # ---------------------------------------------------------
    # DATASET OVERVIEW
    # ---------------------------------------------------------
    s = r['summary']

    story += [
        _pdf_table([
            ['Metric', 'Value'],
            ['Total Rows', f"{s['rows']:,}"],
            ['Total Columns', f"{s['columns']:,}"],
            [
                'Missing Cells',
                f"{s['missing_cells']:,} ({s['missing_percent']:.2f}%)"
            ],
            [
                'Duplicate Rows',
                f"{s['duplicate_rows']:,} ({s['duplicate_percent']:.2f}%)"
            ]
        ]),
        Spacer(1, 12),
        Paragraph('Column Profile', styles['Heading2'])
    ]

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    cp = [
        [
            'Column',
            'Data Type',
            'Semantic Type',
            'Role',
            'Missing',
            'Missing %',
            'Unique'
        ]
    ]

    cp += [
        [
            c['name'],
            c['dtype'],
            str(c.get('semantic_type', '')).replace('_', ' ').title(),
            str(c.get('role', '')).replace('_', ' ').title(),
            str(c['missing']),
            f"{c['missing_percent']:.2f}%",
            str(c['unique'])
        ]
        for c in r['columns']
    ]

    story.append(_pdf_table(cp))

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            'Numerical Statistics',
            styles['Heading2']
        )
    )

    st = [
        ['Column', 'Min', 'Max', 'Mean', 'Median', 'Std Dev']
    ]

    st += [
        [
            x['name'],
            str(x['min']),
            str(x['max']),
            str(x['mean']),
            str(x['median']),
            str(x['std'])
        ]
        for x in r['statistics']
    ]

    story.append(_pdf_table(st))

    # ---------------------------------------------------------
    # DATA QUALITY SUMMARY
    # ---------------------------------------------------------
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            'Data Quality Summary',
            styles['Heading2']
        )
    )

    for f in r['quality']['findings']:
        icon = '✓ ' if f['type'] == 'good' else '! '

        story.append(
            Paragraph(
                icon + f['text'],
                styles['Small']
            )
        )

        story.append(Spacer(1, 3))

    # ---------------------------------------------------------
    # DUPLICATE GROUPS / AFFECTED RECORDS
    # ---------------------------------------------------------
    duplicate_groups = r.get('duplicate_groups', [])
    if duplicate_groups:
        story.append(Spacer(1, 12))
        story.append(Paragraph('Duplicate Groups', styles['Heading2']))
        for group in duplicate_groups:
            story.append(Paragraph(
                f"Group {group['group_id']}: {group['count']} matching records; "
                f"{group['duplicate_count']} duplicate occurrence(s).",
                styles['Small']
            ))
            group_table = [['Source Row', *[str(c) for c in r['columns'] and r.get('preview', {}).get('columns', [])]]]
            rows = group.get('rows', [])
            columns = r.get('preview', {}).get('columns', [])
            group_table = [['Source Row', *columns]]
            for row in rows:
                values = row.get('values', {})
                group_table.append([
                    row.get('row_number', ''),
                    *[_pdf_value(values.get(c)) for c in columns]
                ])
            if len(group_table) > 1:
                story.append(_pdf_table(group_table))
                story.append(Spacer(1, 7))
    else:
        story.append(Spacer(1, 12))
        story.append(Paragraph('Duplicate Groups', styles['Heading2']))
        story.append(Paragraph('No duplicate groups detected.', styles['Small']))

    inconsistent_details = r.get('type_inconsistencies', [])
    if inconsistent_details:
        story.append(Spacer(1, 12))
        story.append(Paragraph('Type Inconsistencies', styles['Heading2']))
        ti = [['Column', 'Issue', 'Details']]
        for item in inconsistent_details:
            ti.append([
                item.get('name', ''),
                item.get('type', '').replace('_', ' ').title(),
                item.get('summary', item.get('message', ''))
            ])
        story.append(_pdf_table(ti))

    outliers = r.get('report', {}).get('outliers', [])
    if outliers:
        story.append(Spacer(1, 12))
        story.append(Paragraph('Potential Outliers', styles['Heading2']))
        story.append(_pdf_table([
            ['Column', 'Outlier Count'],
            *[[x.get('name', ''), x.get('count', 0)] for x in outliers]
        ]))

    # ---------------------------------------------------------
    # VISUAL ANALYSIS
    # ---------------------------------------------------------
    story.append(PageBreak())

    story.append(
        Paragraph(
            'Visual Analysis',
            styles['Heading1']
        )
    )

    # IMPORTANT:
    # analysis.py produces Plotly histogram type = "histogram".
    # The old code checked only "hist", causing those 3 graphs
    # to fail and be silently skipped.
    for c in r['charts']:
        try:
            chart_data = c.get('data', [])

            if not chart_data:
                continue

            chart = chart_data[0]
            chart_type = chart.get('type')

            if chart_type == 'bar' and chart.get('orientation') == 'h':
                # Horizontal Top-N chart: labels are on y, counts on x.
                img = _chart_png(
                    c['title'],
                    'barh',
                    chart.get('y', []),
                    chart.get('x', [])
                )

            elif chart_type == 'bar':
                img = _chart_png(
                    c['title'],
                    'bar',
                    chart.get('x', []),
                    chart.get('y', [])
                )

            elif chart_type == 'histogram':
                img = _chart_png(
                    c['title'],
                    'hist',
                    [],
                    chart.get('x', [])
                )

            elif chart_type == 'box':
                img = _chart_png(
                    c['title'],
                    'box',
                    [],
                    chart.get('y', [])
                )

            else:
                continue

            # KeepTogether: the chart title never ends up alone at the
            # bottom of a page while its chart starts on the next one.
            story.append(
                KeepTogether([
                    Paragraph(
                        c['title'],
                        styles['Heading3']
                    ),
                    Image(
                        img,
                        width=175 * mm,
                        height=80 * mm
                    ),
                    Spacer(1, 7)
                ])
            )

        except Exception as e:
            print(
                f"Could not add chart "
                f"'{c.get('title', 'Unknown')}': {e}"
            )

    doc.build(story)

    return bio.getvalue()


# =============================================================
# EXCEL REPORT
# =============================================================

def _style_ws(ws):
    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color='FFFFFF'
        )
        cell.fill = PatternFill(
            'solid',
            fgColor='081B3D'
        )
        cell.alignment = Alignment(
            horizontal='center'
        )

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for col in ws.columns:
        letter = get_column_letter(
            col[0].column
        )

        maxlen = max(
            len(str(c.value or ''))
            for c in col[:100]
        )

        ws.column_dimensions[letter].width = min(
            max(maxlen + 2, 12),
            38
        )


def build_excel(r, name):
    bio = io.BytesIO()

    wb = Workbook()

    # ---------------------------------------------------------
    # OVERVIEW
    # ---------------------------------------------------------
    ws = wb.active
    ws.title = 'Overview'

    s = r['summary']

    rows = [
        ['Metric', 'Value'],
        ['File Name', name],
        ['Total Rows', s['rows']],
        ['Total Columns', s['columns']],
        ['Missing Cells', s['missing_cells']],
        ['Missing %', s['missing_percent'] / 100],
        ['Duplicate Rows', s['duplicate_rows']],
        ['Duplicate %', s['duplicate_percent'] / 100]
    ]

    for row in rows:
        ws.append(row)

    _style_ws(ws)

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    ws2 = wb.create_sheet('Column Profile')

    ws2.append([
        'Column Name',
        'Data Type',
        'Semantic Type',
        'Role',
        'Missing Values',
        'Missing %',
        'Unique Values'
    ])

    for c in r['columns']:
        ws2.append([
            c['name'],
            c['dtype'],
            c.get('semantic_type', ''),
            c.get('role', ''),
            c['missing'],
            c['missing_percent'] / 100,
            c['unique']
        ])

    _style_ws(ws2)

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    ws3 = wb.create_sheet('Numerical Statistics')

    ws3.append([
        'Column',
        'Minimum',
        'Maximum',
        'Mean',
        'Median',
        'Std Dev'
    ])

    for x in r['statistics']:
        ws3.append([
            x['name'],
            x['min'],
            x['max'],
            x['mean'],
            x['median'],
            x['std']
        ])

    _style_ws(ws3)

    # ---------------------------------------------------------
    # MISSING VALUES
    # ---------------------------------------------------------
    ws4 = wb.create_sheet('Missing Values')

    ws4.append([
        'Column',
        'Missing Values',
        'Missing %'
    ])

    for x in r['missing']:
        ws4.append([
            x['name'],
            x['count'],
            x['percent'] / 100
        ])

    _style_ws(ws4)

    # ---------------------------------------------------------
    # QUALITY SUMMARY
    # ---------------------------------------------------------
    ws5 = wb.create_sheet('Quality Summary')

    ws5.append([
        'Type',
        'Finding'
    ])

    for f in r['quality']['findings']:
        ws5.append([
            f['type'],
            f['text']
        ])

    _style_ws(ws5)

    # ---------------------------------------------------------
    # DUPLICATE ANALYSIS
    # ---------------------------------------------------------
    ws6 = wb.create_sheet('Duplicate Analysis')

    ws6.append([
        'Metric',
        'Value'
    ])

    ws6.append([
        'Duplicate Rows',
        s['duplicate_rows']
    ])

    ws6.append([
        'Duplicate %',
        s['duplicate_percent'] / 100
    ])

    _style_ws(ws6)

    # Target sheet order: ... Missing Values, Duplicate Analysis,
    # Quality Summary.
    wb.move_sheet('Duplicate Analysis', offset=-1)

    # ---------------------------------------------------------
    # DUPLICATE GROUP DETAILS
    # ---------------------------------------------------------
    ws7 = wb.create_sheet('Duplicate Groups')
    columns = r.get('preview', {}).get('columns', [])
    ws7.append(['Group ID', 'Source Row', 'Duplicate Group Count', 'Duplicate Occurrences', *columns])
    for group in r.get('duplicate_groups', []):
        for row in group.get('rows', []):
            values = row.get('values', {})
            ws7.append([
                group.get('group_id', ''),
                row.get('row_number', ''),
                group.get('count', 0),
                group.get('duplicate_count', 0),
                *[values.get(c) for c in columns]
            ])
    _style_ws(ws7)

    # ---------------------------------------------------------
    # TYPE INCONSISTENCIES
    # ---------------------------------------------------------
    ws8 = wb.create_sheet('Type Inconsistencies')
    ws8.append(['Column', 'Issue Type', 'Invalid Count', 'Formatted Count', 'Summary'])
    for item in r.get('type_inconsistencies', []):
        ws8.append([
            item.get('name', ''),
            item.get('type', ''),
            item.get('invalid_count', 0),
            item.get('formatted_count', 0),
            item.get('summary', item.get('message', ''))
        ])
    _style_ws(ws8)

    # ---------------------------------------------------------
    # POTENTIAL OUTLIERS
    # ---------------------------------------------------------
    ws9 = wb.create_sheet('Potential Outliers')
    ws9.append(['Column', 'Outlier Count'])
    for item in r.get('report', {}).get('outliers', []):
        ws9.append([item.get('name', ''), item.get('count', 0)])
    _style_ws(ws9)

    # ---------------------------------------------------------
    # PERCENTAGE FORMATTING
    # ---------------------------------------------------------
    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for cell in row:

                header = str(
                    sh.cell(
                        1,
                        cell.column
                    ).value or ''
                )

                if (
                    isinstance(cell.value, float)
                    and 0 <= cell.value <= 1
                    and (
                        '%' in header
                        or 'Percent' in header
                    )
                ):
                    cell.number_format = '0.00%'

    wb.save(bio)

    return bio.getvalue()


# =============================================================
# CSV REPORT
# =============================================================

# =============================================================
# CSV REPORT
# =============================================================


def build_csv(r, name):

    rows = [
        ['Section', 'Metric', 'Value'],
        ['Overview', 'File Name', name],
        [
            'Overview',
            'Total Rows',
            r['summary']['rows']
        ],
        [
            'Overview',
            'Total Columns',
            r['summary']['columns']
        ],
        [
            'Overview',
            'Missing Cells',
            r['summary']['missing_cells']
        ],
        [
            'Overview',
            'Missing %',
            r['summary']['missing_percent']
        ],
        [
            'Overview',
            'Duplicate Rows',
            r['summary']['duplicate_rows']
        ],
        [
            'Overview',
            'Duplicate %',
            r['summary']['duplicate_percent']
        ]
    ]

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    for c in r['columns']:
        rows.append([
            'Column Profile',
            c['name'],
            (
                f"type={c['dtype']}; "
                f"semantic_type={c.get('semantic_type', '')}; "
                f"role={c.get('role', '')}; "
                f"missing={c['missing']}; "
                f"missing_pct={c['missing_percent']:.2f}; "
                f"unique={c['unique']}"
            )
        ])

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    for x in r['statistics']:
        rows.append([
            'Numerical Statistics',
            x['name'],
            (
                f"min={x['min']}; "
                f"max={x['max']}; "
                f"mean={x['mean']}; "
                f"median={x['median']}; "
                f"std={x['std']}"
            )
        ])

    # ---------------------------------------------------------
    # MISSING VALUES
    # ---------------------------------------------------------
    for x in r['missing']:
        rows.append([
            'Missing Values',
            x['name'],
            (
                f"missing={x['count']}; "
                f"missing_pct={x['percent']:.2f}"
            )
        ])

    # ---------------------------------------------------------
    # DUPLICATE ANALYSIS
    # ---------------------------------------------------------
    rows.append([
        'Duplicate Analysis',
        'Duplicate Rows',
        r['summary']['duplicate_rows']
    ])

    rows.append([
        'Duplicate Analysis',
        'Duplicate %',
        r['summary']['duplicate_percent']
    ])

    # ---------------------------------------------------------
    # QUALITY SUMMARY
    # ---------------------------------------------------------
    for f in r['quality']['findings']:
        rows.append([
            'Quality Summary',
            f['type'],
            f['text']
        ])

    # ---------------------------------------------------------
    # DUPLICATE GROUP DETAILS
    # ---------------------------------------------------------
    for group in r.get('duplicate_groups', []):
        for row in group.get('rows', []):
            values = row.get('values', {})
            rows.append([
                'Duplicate Groups',
                f"Group {group.get('group_id', '')} — Row {row.get('row_number', '')}",
                (
                    f"group_count={group.get('count', 0)}; "
                    f"duplicate_occurrences={group.get('duplicate_count', 0)}; "
                    + '; '.join(f"{c}={values.get(c, '')}" for c in r.get('preview', {}).get('columns', []))
                )
            ])

    # ---------------------------------------------------------
    # TYPE INCONSISTENCIES
    # ---------------------------------------------------------
    for item in r.get('type_inconsistencies', []):
        rows.append([
            'Type Inconsistencies',
            item.get('name', ''),
            item.get('summary', item.get('message', ''))
        ])

    # ---------------------------------------------------------
    # POTENTIAL OUTLIERS
    # ---------------------------------------------------------
    for item in r.get('report', {}).get('outliers', []):
        rows.append([
            'Potential Outliers',
            item.get('name', ''),
            item.get('count', 0)
        ])

    out = io.StringIO()

    pd.DataFrame(
        rows[1:],
        columns=rows[0]
    ).to_csv(
        out,
        index=False
    )

    return out.getvalue().encode('utf-8-sig')