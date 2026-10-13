import io
from xml.sax.saxutils import escape
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

PURPLE = '#6744ff'
NAVY = '#081b3d'
TEXT = '#1b2b4a'


def _duplicate_summary_rows(s):
    """Duplicate figures straight from the analysis summary (same numbers as the dashboard)."""
    ignoring = s.get('duplicate_rows_ignoring_identifiers', s['duplicate_rows'])
    ignoring_pct = s.get('duplicate_percent_ignoring_identifiers', s['duplicate_percent'])
    ignored = ', '.join(s.get('duplicate_ignored_identifier_columns') or []) or 'none'

    return [
        ('Duplicate Rows (exact: all columns identical)', s['duplicate_rows']),
        ('Duplicate % (exact)', s['duplicate_percent']),
        ('Exact Duplicate Groups', s.get('exact_duplicate_groups', 0)),
        ('Duplicate Rows (ignoring identifier columns)', ignoring),
        ('Duplicate % (ignoring identifier columns)', ignoring_pct),
        ('Duplicate Groups (total)', s.get('duplicate_groups_total', 0)),
        ('Identifier Columns Ignored', ignored),
    ]


def _duplicate_group_rows(r):
    """One row per duplicate group: id, type, original row, all rows, identifier columns that differ."""
    label = {'exact': 'Exact', 'identifier_excluded': 'Identifiers differ'}

    return [
        [
            g['group_id'],
            label.get(g.get('match_type'), g.get('match_type', '')),
            g.get('original_row', g['row_numbers'][0] if g['row_numbers'] else ''),
            ', '.join(str(x) for x in g['row_numbers']),
            ', '.join(g.get('differing_identifier_columns') or [])
        ]
        for g in r.get('duplicate_groups', [])
    ]


def _pdf_table(data, widths=None):
    t = Table(data, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(NAVY)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 7.5),
        ('GRID', (0, 0), (-1, -1), .25, colors.HexColor('#b9c6dc')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [
            colors.white,
            colors.HexColor('#f4f7fc')
        ]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5)
    ]))
    return t


def _chart_png(title, kind, labels, values, box_stats=None):
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    if kind == 'bar':
        labels = [str(label) for label in labels]
        positions = list(range(len(labels)))

        ax.bar(positions, values, color='#6744ff')

        # Keep category / date labels readable: rotate long or many
        # labels and show only every n-th one when there are many bars.
        step = max(1, int(np.ceil(len(labels) / 12)))
        longest = max((len(label) for label in labels), default=0)
        rotate = len(labels) > 6 or longest > 10

        ax.set_xticks(positions)
        ax.set_xticklabels(
            [
                label if i % step == 0 else ''
                for i, label in enumerate(labels)
            ],
            rotation=45 if rotate else 0,
            ha='right' if rotate else 'center',
            fontsize=8
        )

    elif kind == 'barh':
        labels = [str(label) for label in labels]

        ax.barh(labels, values, color='#6744ff')
        ax.invert_yaxis()
        ax.tick_params(axis='y', labelsize=8)

    elif kind == 'hist':
        values = np.asarray(values, dtype=float)

        if len(values):
            bins = min(
                15,
                max(5, int(np.sqrt(max(1, len(values)))))
            )
            ax.hist(values, bins=bins, color='#6744ff')

    elif kind == 'box':
        data = np.asarray(values, dtype=float)

        if box_stats and box_stats.get('q1') is not None:
            # Draw exactly the box the analysis computed, so the outlier
            # points in the PDF are the ones counted in the analysis
            # (matplotlib's own quartiles can differ from the dashboard's).
            low = box_stats['lower_whisker']
            high = box_stats['upper_whisker']

            ax.bxp(
                [{
                    'med': box_stats['median'],
                    'q1': box_stats['q1'],
                    'q3': box_stats['q3'],
                    'whislo': low,
                    'whishi': high,
                    'fliers': data[(data < low) | (data > high)]
                }],
                vert=True,
                patch_artist=True,
                boxprops={'facecolor': '#6744ff'}
            )

        else:
            # Fewer than 4 numeric values (outliers not evaluated): box only.
            ax.boxplot(
                data,
                vert=True,
                patch_artist=True,
                showfliers=False,
                boxprops={'facecolor': '#6744ff'}
            )

    ax.set_title(title)
    ax.grid(axis='x' if kind == 'barh' else 'y', alpha=.18)

    fig.tight_layout()

    bio = io.BytesIO()
    fig.savefig(
        bio,
        format='png',
        dpi=150,
        bbox_inches='tight'
    )

    plt.close(fig)
    bio.seek(0)

    return bio


def build_pdf(r, name):
    bio = io.BytesIO()

    doc = SimpleDocTemplate(
        bio,
        pagesize=A4,
        rightMargin=13 * mm,
        leftMargin=13 * mm,
        topMargin=13 * mm,
        bottomMargin=13 * mm
    )

    styles = getSampleStyleSheet()

    styles.add(
        ParagraphStyle(
            name='Small',
            parent=styles['BodyText'],
            fontSize=8.5,
            leading=11,
            textColor=colors.HexColor('#40516d')
        )
    )

    # ---------------------------------------------------------
    # REPORT HEADER
    # ---------------------------------------------------------
    story = [
        Paragraph(
            'DataLens — Data Quality Analysis Report',
            styles['Title']
        ),
        Paragraph(
            f'File: {name}',
            styles['BodyText']
        ),
        Spacer(1, 8)
    ]

    # ---------------------------------------------------------
    # DATASET OVERVIEW
    # ---------------------------------------------------------
    s = r['summary']

    story += [
        _pdf_table([
            ['Metric', 'Value'],
            ['Total Rows', f"{s['rows']:,}"],
            ['Total Columns', f"{s['columns']:,}"],
            [
                'Missing Cells',
                f"{s['missing_cells']:,} ({s['missing_percent']:.2f}%)"
            ],
            [
                'Duplicate Rows (exact)',
                f"{s['duplicate_rows']:,} ({s['duplicate_percent']:.2f}%)"
            ],
            [
                'Duplicates ignoring identifiers',
                f"{s.get('duplicate_rows_ignoring_identifiers', s['duplicate_rows']):,} "
                f"({s.get('duplicate_percent_ignoring_identifiers', s['duplicate_percent']):.2f}%)"
            ]
        ]),
        Spacer(1, 12),
        Paragraph('Column Profile', styles['Heading2'])
    ]

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    cp = [
        [
            'Column',
            'Data Type',
            'Semantic Type',
            'Role',
            'Missing',
            'Missing %',
            'Unique'
        ]
    ]

    cp += [
        [
            c['name'],
            c['dtype'],
            str(c.get('semantic_type', '')).replace('_', ' ').title(),
            str(c.get('role', '')).replace('_', ' ').title(),
            str(c['missing']),
            f"{c['missing_percent']:.2f}%",
            str(c['unique'])
        ]
        for c in r['columns']
    ]

    story.append(_pdf_table(cp))

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            'Numerical Statistics',
            styles['Heading2']
        )
    )

    st = [
        ['Column', 'Min', 'Max', 'Mean', 'Median', 'Std Dev']
    ]

    st += [
        [
            x['name'],
            str(x['min']),
            str(x['max']),
            str(x['mean']),
            str(x['median']),
            str(x['std'])
        ]
        for x in r['statistics']
    ]

    story.append(_pdf_table(st))

    # ---------------------------------------------------------
    # DATA QUALITY SUMMARY
    # ---------------------------------------------------------
    story.append(Spacer(1, 12))
    story.append(
        Paragraph(
            'Data Quality Summary',
            styles['Heading2']
        )
    )

    for f in r['quality']['findings']:
        icon = '✓ ' if f['type'] == 'good' else '! '

        story.append(
            Paragraph(
                icon + escape(f['text']),
                styles['Small']
            )
        )

        story.append(Spacer(1, 3))

    # ---------------------------------------------------------
    # VISUAL ANALYSIS
    # ---------------------------------------------------------
    story.append(PageBreak())

    story.append(
        Paragraph(
            'Visual Analysis',
            styles['Heading1']
        )
    )

    # IMPORTANT:
    # analysis.py produces Plotly histogram type = "histogram".
    # The old code checked only "hist", causing those 3 graphs
    # to fail and be silently skipped.
    for c in r['charts']:
        try:
            chart_data = c.get('data', [])

            if not chart_data:
                continue

            chart = chart_data[0]
            chart_type = chart.get('type')

            if chart_type == 'bar' and chart.get('orientation') == 'h':
                # Horizontal Top-N chart: labels are on y, counts on x.
                img = _chart_png(
                    c['title'],
                    'barh',
                    chart.get('y', []),
                    chart.get('x', [])
                )

            elif chart_type == 'bar':
                img = _chart_png(
                    c['title'],
                    'bar',
                    chart.get('x', []),
                    chart.get('y', [])
                )

            elif chart_type == 'histogram':
                img = _chart_png(
                    c['title'],
                    'hist',
                    [],
                    chart.get('x', [])
                )

            elif chart_type == 'box':
                # Box plots carry their values in `x` when horizontal.
                img = _chart_png(
                    c['title'],
                    'box',
                    [],
                    chart.get('x', []) if chart.get('orientation') == 'h' else chart.get('y', []),
                    box_stats=c.get('box_stats')
                )

            else:
                continue

            # KeepTogether: the chart title never ends up alone at the
            # bottom of a page while its chart starts on the next one.
            unit = [
                Paragraph(
                    c['title'],
                    styles['Heading3']
                ),
                Image(
                    img,
                    width=175 * mm,
                    height=80 * mm
                )
            ]

            if c.get('note'):
                unit.append(
                    Paragraph(
                        escape(str(c['note'])),
                        styles['Small']
                    )
                )

            unit.append(Spacer(1, 7))

            story.append(KeepTogether(unit))

        except Exception as e:
            print(
                f"Could not add chart "
                f"'{c.get('title', 'Unknown')}': {e}"
            )

    doc.build(story)

    return bio.getvalue()


# =============================================================
# EXCEL REPORT
# =============================================================

def _style_ws(ws):
    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color='FFFFFF'
        )
        cell.fill = PatternFill(
            'solid',
            fgColor='081B3D'
        )
        cell.alignment = Alignment(
            horizontal='center'
        )

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions

    for col in ws.columns:
        letter = get_column_letter(
            col[0].column
        )

        maxlen = max(
            len(str(c.value or ''))
            for c in col[:100]
        )

        ws.column_dimensions[letter].width = min(
            max(maxlen + 2, 12),
            38
        )


def build_excel(r, name):
    bio = io.BytesIO()

    wb = Workbook()

    # ---------------------------------------------------------
    # OVERVIEW
    # ---------------------------------------------------------
    ws = wb.active
    ws.title = 'Overview'

    s = r['summary']

    rows = [
        ['Metric', 'Value'],
        ['File Name', name],
        ['Total Rows', s['rows']],
        ['Total Columns', s['columns']],
        ['Missing Cells', s['missing_cells']],
        ['Missing %', s['missing_percent'] / 100],
        ['Duplicate Rows', s['duplicate_rows']],
        ['Duplicate %', s['duplicate_percent'] / 100],
        ['Duplicate Rows (ignoring identifiers)', s.get('duplicate_rows_ignoring_identifiers', s['duplicate_rows'])],
        ['Duplicate % (ignoring identifiers)', s.get('duplicate_percent_ignoring_identifiers', s['duplicate_percent']) / 100]
    ]

    for row in rows:
        ws.append(row)

    _style_ws(ws)

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    ws2 = wb.create_sheet('Column Profile')

    ws2.append([
        'Column Name',
        'Data Type',
        'Semantic Type',
        'Role',
        'Missing Values',
        'Missing %',
        'Unique Values'
    ])

    for c in r['columns']:
        ws2.append([
            c['name'],
            c['dtype'],
            c.get('semantic_type', ''),
            c.get('role', ''),
            c['missing'],
            c['missing_percent'] / 100,
            c['unique']
        ])

    _style_ws(ws2)

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    ws3 = wb.create_sheet('Numerical Statistics')

    ws3.append([
        'Column',
        'Minimum',
        'Maximum',
        'Mean',
        'Median',
        'Std Dev'
    ])

    for x in r['statistics']:
        ws3.append([
            x['name'],
            x['min'],
            x['max'],
            x['mean'],
            x['median'],
            x['std']
        ])

    _style_ws(ws3)

    # ---------------------------------------------------------
    # MISSING VALUES
    # ---------------------------------------------------------
    ws4 = wb.create_sheet('Missing Values')

    ws4.append([
        'Column',
        'Missing Values',
        'Missing %'
    ])

    for x in r['missing']:
        ws4.append([
            x['name'],
            x['count'],
            x['percent'] / 100
        ])

    _style_ws(ws4)

    # ---------------------------------------------------------
    # QUALITY SUMMARY
    # ---------------------------------------------------------
    ws5 = wb.create_sheet('Quality Summary')

    ws5.append([
        'Type',
        'Finding'
    ])

    for f in r['quality']['findings']:
        ws5.append([
            f['type'],
            f['text']
        ])

    _style_ws(ws5)

    # ---------------------------------------------------------
    # DUPLICATE ANALYSIS
    # ---------------------------------------------------------
    ws6 = wb.create_sheet('Duplicate Analysis')

    ws6.append([
        'Metric',
        'Value'
    ])

    for label, value in _duplicate_summary_rows(s):
        ws6.append([
            label,
            value / 100 if '%' in label else value
        ])

    _style_ws(ws6)

    group_rows = _duplicate_group_rows(r)

    if group_rows:
        ws6.append([])
        ws6.append([
            'Group',
            'Match Type',
            'Original Row',
            'Rows (source file)',
            'Identifier Columns That Differ'
        ])

        for row in group_rows:
            ws6.append(row)

    wb.move_sheet('Duplicate Analysis', offset=-1)

    # ---------------------------------------------------------
    # PERCENTAGE FORMATTING
    # ---------------------------------------------------------
    for sh in wb.worksheets:
        for row in sh.iter_rows(min_row=2):
            for cell in row:
                header = str(
                    sh.cell(1, cell.column).value or ''
                )

                row_label = str(
                    sh.cell(cell.row, 1).value or ''
                )

                is_percentage = (
                    '%' in header
                    or 'Percent' in header
                    or '%' in row_label
                    or 'Percent' in row_label
                )

                if (
                    isinstance(cell.value, (int, float))
                    and not isinstance(cell.value, bool)
                    and 0 <= cell.value <= 1
                    and is_percentage
                ):
                    cell.number_format = '0.00%'

    wb.save(bio)

    return bio.getvalue()


# =============================================================
# CSV REPORT
# =============================================================

# =============================================================
# CSV REPORT
# =============================================================


def build_csv(r, name):

    rows = [
        ['Section', 'Metric', 'Value'],
        ['Overview', 'File Name', name],
        [
            'Overview',
            'Total Rows',
            r['summary']['rows']
        ],
        [
            'Overview',
            'Total Columns',
            r['summary']['columns']
        ],
        [
            'Overview',
            'Missing Cells',
            r['summary']['missing_cells']
        ],
        [
            'Overview',
            'Missing %',
            f"{r['summary']['missing_percent']:.2f}%"
        ],
        [
            'Overview',
            'Duplicate Rows',
            r['summary']['duplicate_rows']
        ],
        [
            'Overview',
            'Duplicate %',
            f"{r['summary']['duplicate_percent']:.2f}%"
        ],
        [
            'Overview',
            'Duplicate Rows (ignoring identifiers)',
            r['summary'].get('duplicate_rows_ignoring_identifiers', r['summary']['duplicate_rows'])
        ],
        [
            'Overview',
            'Duplicate % (ignoring identifiers)',
            (
                f"{r['summary'].get('duplicate_percent_ignoring_identifiers', r['summary']['duplicate_percent']):.2f}%"
            )
        ]
    ]

    # ---------------------------------------------------------
    # COLUMN PROFILE
    # ---------------------------------------------------------
    for c in r['columns']:
        rows.append([
            'Column Profile',
            c['name'],
            (
                f"type={c['dtype']}; "
                f"semantic_type={c.get('semantic_type', '')}; "
                f"role={c.get('role', '')}; "
                f"missing={c['missing']}; "
                f"missing_pct={c['missing_percent']:.2f}%; "
                f"unique={c['unique']}"
            )
        ])

    # ---------------------------------------------------------
    # NUMERICAL STATISTICS
    # ---------------------------------------------------------
    for x in r['statistics']:
        rows.append([
            'Numerical Statistics',
            x['name'],
            (
                f"min={x['min']}; "
                f"max={x['max']}; "
                f"mean={x['mean']}; "
                f"median={x['median']}; "
                f"std={x['std']}"
            )
        ])

    # ---------------------------------------------------------
    # MISSING VALUES
    # ---------------------------------------------------------
    for x in r['missing']:
        rows.append([
            'Missing Values',
            x['name'],
            (
                f"missing={x['count']}; "
                f"missing_pct={x['percent']:.2f}%"
            )
        ])

    # ---------------------------------------------------------
    # DUPLICATE ANALYSIS
    # ---------------------------------------------------------
    for label, value in _duplicate_summary_rows(r['summary']):
        if '%' in label and isinstance(value, (int, float)):
            value = f"{value:.2f}%"

        rows.append([
            'Duplicate Analysis',
            label,
            value
        ])

    for g in _duplicate_group_rows(r):
        rows.append([
            'Duplicate Analysis',
            f'Group {g[0]}',
            (
                f'type={g[1]}; original_row={g[2]}; rows={g[3]}'
                + (f'; identifier_columns_that_differ={g[4]}' if g[4] else '')
            )
        ])

    # ---------------------------------------------------------
    # QUALITY SUMMARY
    # ---------------------------------------------------------
    for f in r['quality']['findings']:
        rows.append([
            'Quality Summary',
            f['type'],
            f['text']
        ])

    out = io.StringIO()

    pd.DataFrame(
        rows[1:],
        columns=rows[0]
    ).to_csv(
        out,
        index=False
    )

    return out.getvalue().encode('utf-8-sig')